"""In-memory TXT/PDF/DOCX extraction and integration tests (no model download)."""

from __future__ import annotations

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from src.core.services.summarization.document_extraction_service import (
    CorruptedDocumentError,
    DocumentExtractionService,
    DocumentLimitError,
    EmptyDocumentError,
    InvalidDocumentError,
    InvalidTextEncodingError,
    NoExtractableTextError,
    PasswordProtectedPdfError,
    UnsupportedDocumentTypeError,
)
from src.infrastructure.di.bootstrap import bootstrap_di
from src.infrastructure.di.inject import resolve
from src.infrastructure.ml.mt5_model_provider import Mt5ModelProvider


@pytest.fixture
def extractor() -> DocumentExtractionService:
    return DocumentExtractionService()


def pdf_with_text(*pages: str) -> bytes:
    """Make a text-layer PDF with a deterministic, custom Unicode CMap.

    Synthetic font glyphs are used so this test never requires an installed
    Persian font. The ToUnicode map contains real Persian codepoints; test
    production PDFs with real embedded fonts separately on the target system.
    """
    writer = PdfWriter()
    for text in pages:
        page = writer.add_blank_page(width=612, height=792)
        if not text:
            continue
        letters = list(dict.fromkeys(text))
        assert len(letters) < 256
        mapping = {char: i + 1 for i, char in enumerate(letters)}
        mappings = "\n".join(f"<{index:02X}> <{ord(char):04X}>" for char, index in mapping.items())
        cmap = (
            "/CIDInit /ProcSet findresource begin\n"
            "12 dict begin\nbegincmap\n"
            "/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def\n"
            "/CMapName /Adobe-Identity-UCS def\n/CMapType 2 def\n"
            "1 begincodespacerange\n<00> <FF>\nendcodespacerange\n"
            f"{len(mapping)} beginbfchar\n{mappings}\nendbfchar\n"
            "endcmap\nCMapName currentdict /CMap defineresource pop\nend\nend"
        ).encode("ascii")
        cmap_stream = DecodedStreamObject()
        cmap_stream.set_data(cmap)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
                NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
                NameObject("/ToUnicode"): writer._add_object(cmap_stream),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        hex_text = "".join(f"{mapping[char]:02X}" for char in text)
        content_stream = DecodedStreamObject()
        content_stream.set_data(f"BT /F1 12 Tf 50 700 Td <{hex_text}> Tj ET".encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(content_stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def empty_pdf(*, pages: int = 1, encrypted: bool = False) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=612, height=792)
    if encrypted:
        writer.encrypt(user_password="private-password")
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def docx_bytes(*paragraphs: str, table: bool = False) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    if table:
        grid = doc.add_table(rows=2, cols=2)
        grid.cell(0, 0).text = "نام"
        grid.cell(0, 1).text = "مقدار"
        grid.cell(1, 0).text = "کتاب"
        grid.cell(1, 1).text = "۲"
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_smoke_txt_extraction(extractor: DocumentExtractionService) -> None:
    actual = extractor.extract("این یک متن فارسی است.".encode(), "sample.TXT")
    assert actual == "این یک متن فارسی است."
    print("TXT:", actual)


def test_txt_utf8_bom(extractor: DocumentExtractionService) -> None:
    assert extractor.extract(b"\xef\xbb\xbf" + "سلام".encode(), "a.txt") == "سلام"


def test_txt_newline_and_unicode_preservation(extractor: DocumentExtractionService) -> None:
    text = "\ufeff  می‌روم، ۱۲۳!\r\n\r\nعربی: ي، ك \r"
    assert extractor.extract(text.encode("utf-8"), "a.txt") == "می‌روم، ۱۲۳!\n\nعربی: ي، ك"


@pytest.mark.parametrize("content", [b"", "".encode()])
def test_empty_txt_bytes(extractor: DocumentExtractionService, content: bytes) -> None:
    with pytest.raises(EmptyDocumentError):
        extractor.extract(content, "empty.txt")


@pytest.mark.parametrize("content", [b" \t\r\n ", "\ufeff\n".encode("utf-8")])
def test_whitespace_txt(extractor: DocumentExtractionService, content: bytes) -> None:
    with pytest.raises(NoExtractableTextError):
        extractor.extract_txt(content)


def test_invalid_txt_encoding(extractor: DocumentExtractionService) -> None:
    with pytest.raises(InvalidTextEncodingError):
        extractor.extract_txt(b"\xff\xfe\x80")


@pytest.mark.parametrize("content", [b"a\x00b", b"a\x01b", b"a\x7fb"])
def test_binary_txt_rejected(extractor: DocumentExtractionService, content: bytes) -> None:
    with pytest.raises(InvalidDocumentError):
        extractor.extract_txt(content)


def test_smoke_pdf_persian_text(extractor: DocumentExtractionService) -> None:
    text = "این یک متن فارسی برای آزمون است"
    actual = extractor.extract(pdf_with_text(text), "sample.PDF")
    assert actual == text
    print("PDF:", actual)


def test_pdf_multiple_pages_and_blank_middle(extractor: DocumentExtractionService) -> None:
    actual = extractor.extract_pdf(pdf_with_text("صفحه اول", "", "صفحه سوم"))
    assert actual == "صفحه اول\n\nصفحه سوم"


def test_pdf_english_text(extractor: DocumentExtractionService) -> None:
    assert extractor.extract_pdf(pdf_with_text("English text")) == "English text"


@pytest.mark.parametrize("data", [empty_pdf(), empty_pdf(pages=0)])
def test_pdf_no_extractable_text(extractor: DocumentExtractionService, data: bytes) -> None:
    with pytest.raises(NoExtractableTextError):
        extractor.extract_pdf(data)


def test_pdf_invalid_signature(extractor: DocumentExtractionService) -> None:
    with pytest.raises(InvalidDocumentError, match="signature"):
        extractor.extract_pdf(b"not a pdf")


def test_pdf_corrupted(extractor: DocumentExtractionService) -> None:
    with pytest.raises(CorruptedDocumentError):
        extractor.extract_pdf(b"%PDF-1.7\nmalformed garbage\n")


def test_pdf_password_protected(extractor: DocumentExtractionService) -> None:
    with pytest.raises(PasswordProtectedPdfError):
        extractor.extract_pdf(empty_pdf(encrypted=True))


def test_pdf_page_limit(extractor: DocumentExtractionService, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(DocumentExtractionService, "MAX_PDF_PAGES", 1)
    with pytest.raises(DocumentLimitError, match="page count"):
        extractor.extract_pdf(empty_pdf(pages=2))


def test_pdf_character_limit(extractor: DocumentExtractionService, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(DocumentExtractionService, "MAX_TEXT_CHARACTERS", 5)
    with pytest.raises(DocumentLimitError, match="character limit"):
        extractor.extract_pdf(pdf_with_text("۱۲۳۴۵۶"))


def test_smoke_docx_persian(extractor: DocumentExtractionService) -> None:
    actual = extractor.extract(docx_bytes("متن فارسی اول", "پاراگراف دوم"), "sample.DOCX")
    assert actual == "متن فارسی اول\nپاراگراف دوم"
    print("DOCX:", actual)


def test_docx_table_document_order(extractor: DocumentExtractionService) -> None:
    doc = Document()
    doc.add_paragraph("مقدمه")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "نام"
    table.cell(0, 1).text = "تعداد"
    doc.add_paragraph("پایان")
    raw = BytesIO()
    doc.save(raw)
    assert extractor.extract_docx(raw.getvalue()) == "مقدمه\nنام\tتعداد\nپایان"


def test_docx_table_and_unicode(extractor: DocumentExtractionService) -> None:
    result = extractor.extract_docx(docx_bytes("می‌روم ۱۲۳", table=True))
    assert "می‌روم ۱۲۳" in result
    assert "کتاب\t۲" in result


def test_docx_empty(extractor: DocumentExtractionService) -> None:
    with pytest.raises(NoExtractableTextError):
        extractor.extract_docx(docx_bytes())


def test_docx_whitespace_only(extractor: DocumentExtractionService) -> None:
    with pytest.raises(NoExtractableTextError):
        extractor.extract_docx(docx_bytes("   ", "\n "))


@pytest.mark.parametrize("data", [b"not-zip", b"PK\x03\x04garbage"])
def test_docx_corrupted(extractor: DocumentExtractionService, data: bytes) -> None:
    with pytest.raises(CorruptedDocumentError):
        extractor.extract_docx(data)


def test_docx_missing_required_xml_parts(extractor: DocumentExtractionService) -> None:
    buffer = BytesIO()
    with ZipFile(buffer, "w", ZIP_DEFLATED) as file:
        file.writestr("test.txt", "invalid content")
    with pytest.raises(CorruptedDocumentError, match="required XML"):
        extractor.extract_docx(buffer.getvalue())


def test_docx_resource_guard(extractor: DocumentExtractionService, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(DocumentExtractionService, "MAX_DOCX_PART_BYTES", 100)
    with pytest.raises(DocumentLimitError, match="oversized ZIP part"):
        extractor.extract_docx(docx_bytes("small text"))


def test_docx_many_members_guard(extractor: DocumentExtractionService, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(DocumentExtractionService, "MAX_DOCX_MEMBERS", 2)
    with pytest.raises(DocumentLimitError, match="too many"):
        extractor.extract_docx(docx_bytes("text"))


def test_docx_character_limit(extractor: DocumentExtractionService, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(DocumentExtractionService, "MAX_TEXT_CHARACTERS", 5)
    with pytest.raises(DocumentLimitError, match="character limit"):
        extractor.extract_docx(docx_bytes("این متن طولانی است"))


@pytest.mark.parametrize("name", ["book.doc", "a.png", "a", "", None, "test.pdf.exe"])
def test_unsupported_extensions(extractor: DocumentExtractionService, name: str | None) -> None:
    with pytest.raises(UnsupportedDocumentTypeError):
        extractor.extract(b"nonempty", name)


@pytest.mark.parametrize("name", ["a.txt", "a.pdf", "a.docx"])
def test_empty_input_all_formats(extractor: DocumentExtractionService, name: str) -> None:
    with pytest.raises(EmptyDocumentError):
        extractor.extract(b"", name)


@pytest.mark.parametrize("value", ["not bytes", None, bytearray(b"hi"), 123])
def test_content_must_be_bytes(extractor: DocumentExtractionService, value: object) -> None:
    with pytest.raises(InvalidDocumentError):
        extractor.extract(value, "file.txt")


def test_raw_input_limit(extractor: DocumentExtractionService, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(DocumentExtractionService, "MAX_INPUT_BYTES", 3)
    with pytest.raises(DocumentLimitError, match="byte limit"):
        extractor.extract_txt(b"four")


def test_di_resolves_without_loading_model() -> None:
    bootstrap_di()
    service = resolve(DocumentExtractionService)
    assert isinstance(service, DocumentExtractionService)
    assert service is not resolve(DocumentExtractionService)


def test_extract_does_not_load_mt5(monkeypatch: pytest.MonkeyPatch) -> None:
    def no_model(*args: object, **kwargs: object) -> None:
        pytest.fail("mT5 was loaded by document extraction")

    monkeypatch.setattr(Mt5ModelProvider, "_ensure_loaded", no_model)
    bootstrap_di()
    service = resolve(DocumentExtractionService)
    assert service.extract("سلام".encode(), "a.txt") == "سلام"


def test_extract_no_temporary_file_writes(
    extractor: DocumentExtractionService, monkeypatch: pytest.MonkeyPatch
) -> None:
    import tempfile

    def no_temp(*args: object, **kwargs: object) -> None:
        pytest.fail("an extraction operation wrote a temporary file")

    monkeypatch.setattr(tempfile, "NamedTemporaryFile", no_temp)
    assert extractor.extract("سلام".encode(), "a.txt") == "سلام"
    assert extractor.extract_pdf(pdf_with_text("سلام")) == "سلام"
    assert extractor.extract_docx(docx_bytes("سلام")) == "سلام"


def test_no_fastapi_dependency_in_core() -> None:
    from pathlib import Path
    import src.core.services.summarization.document_extraction_service as module

    assert "fastapi" not in Path(module.__file__).read_text(encoding="utf-8").lower()
